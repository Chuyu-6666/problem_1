from openpyxl import load_workbook
import numpy as np


def load_attachment_one(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active

    data = np.array([
        row
        for row in worksheet.iter_rows(
            min_row=2,
            min_col=1,
            max_col=3,
            values_only=True,
        )
        if row[0] is not None
    ], dtype=float)

    workbook.close()

    temperature_data = data[:, [0, 1]]
    moisture_data = data[:, [0, 2]]

    return temperature_data, moisture_data
